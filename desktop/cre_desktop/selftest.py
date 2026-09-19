"""`CRE Underwriting --self-test`: exercise the frozen build without a window.

Hits the code paths whose dependencies are hardest to freeze (native
extensions, data files loaded by path): the pro-forma engine, openpyxl
export, python-pptx decks, python-docx memo with matplotlib charts,
reportlab + pdfplumber/pypdfium2, and the Keychain / Credential Manager
backend. Runs against a
throwaway storage folder so the user's deals are never touched.
Prints one line per check; exit code 1 if any check fails.
"""

import json
import os
import shutil
import sys
import tempfile
import traceback
from io import BytesIO
from pathlib import Path

from .osutil import SECRET_STORE_LABEL
from .paths import bundle_root, backend_dir, frontend_dist

# The backend test suite's analytic acquisition fixture (bundled by the spec).
SAMPLE_DEAL = "analytic_acquisition.json"


def _sample_deal() -> dict:
    for candidate in (
        bundle_root() / "selftest" / SAMPLE_DEAL,
        bundle_root() / "backend" / "tests" / "fixtures" / SAMPLE_DEAL,
    ):
        if candidate.exists():
            return json.loads(candidate.read_text())
    raise FileNotFoundError(SAMPLE_DEAL)


def run() -> int:
    scratch = Path(tempfile.mkdtemp(prefix="cre-selftest-"))
    os.environ["CRE_STORAGE_ROOT"] = str(scratch)
    os.environ["MPLCONFIGDIR"] = str(scratch / "mpl")
    os.environ.pop("CRE_ENABLE_BACKUP_SCHEDULER", None)
    sys.path.insert(0, str(backend_dir()))

    failures = 0

    def check(name, fn):
        nonlocal failures
        try:
            detail = fn()
            print(f"PASS  {name}{f' — {detail}' if detail else ''}", flush=True)
        except Exception:  # noqa: BLE001
            failures += 1
            print(f"FAIL  {name}\n{traceback.format_exc()}", flush=True)

    state: dict = {}

    def boot():
        from fastapi.testclient import TestClient

        from app.main import app

        state["client"] = TestClient(app)
        state["deal"] = _sample_deal()
        return None

    def frontend():
        assert (frontend_dist() / "index.html").exists()
        r = state["client"].get("/api/health")
        assert r.status_code == 200, r.text

    def compute():
        r = state["client"].post("/api/compute?detail=true", json={"values": state["deal"]})
        assert r.status_code == 200, r.text
        irr = r.json()["outputs"].get("leveredIrr")
        assert isinstance(irr, (int, float)), irr
        return f"levered IRR {irr:.2%}"

    def excel_export():
        r = state["client"].post("/api/generate/model", json={"values": state["deal"]})
        assert r.status_code == 200, r.text
        import openpyxl

        wb = openpyxl.load_workbook(BytesIO(r.content))
        return f"{len(wb.sheetnames)} sheets"

    def decks():
        r = state["client"].post("/api/deals", json={"name": "Self-test", "inputs": state["deal"]})
        assert r.status_code == 200, r.text
        deal_id = r.json()["id"]
        state["deal_id"] = deal_id
        from pptx import Presentation

        for path in (f"/api/deals/{deal_id}/deck.pptx", f"/api/deals/{deal_id}/ic-deck.pptx"):
            resp = state["client"].get(path)
            assert resp.status_code == 200, f"{path}: {resp.text[:300]}"
            Presentation(BytesIO(resp.content))
        share = state["client"].get(f"/api/deals/{deal_id}/share.html")
        assert share.status_code == 200, share.text[:300]

    def memo():
        r = state["client"].post(
            "/api/scenarios",
            json={"scenarioName": "Self-test", "dealId": state.get("deal_id"), "inputs": state["deal"]},
        )
        assert r.status_code == 200, r.text
        memo_resp = state["client"].post(f"/api/scenarios/{r.json()['id']}/memo", json={})
        assert memo_resp.status_code == 200, memo_resp.text[:300]
        from docx import Document

        doc = Document(BytesIO(memo_resp.content))
        return f"{len(doc.inline_shapes)} embedded chart(s)"

    def pdf_roundtrip():
        from reportlab.pdfgen import canvas

        pdf_path = scratch / "t.pdf"
        c = canvas.Canvas(str(pdf_path))
        c.drawString(72, 720, "Rent Roll Unit 101 1,250")
        c.save()
        import pdfplumber

        with pdfplumber.open(pdf_path) as pdf:
            text = pdf.pages[0].extract_text() or ""
            pdf.pages[0].to_image(resolution=36)  # pypdfium2 native render
        assert "Rent Roll" in text, text

    def keychain():
        import keyring

        from .keys import EXPECTED_BACKENDS

        backend = keyring.get_keyring()
        expected = EXPECTED_BACKENDS.get(sys.platform)
        assert expected is None or type(backend).__module__ == expected, type(backend).__module__
        return type(backend).__module__

    def external_tools():
        from app.services import soffice
        from app.services.extraction import ocr

        return (
            f"LibreOffice {'found' if soffice.is_available() else 'NOT installed'}, "
            f"OCR {'available' if ocr.is_available() else 'NOT installed'} (both optional)"
        )

    check("backend imports + migrations", boot)
    if "client" in state:
        check("frontend bundle + health", frontend)
        check("native compute", compute)
        check("Excel model export (openpyxl)", excel_export)
        check("deck + IC deck + share page (python-pptx, matplotlib)", decks)
        check("IC memo .docx with charts (python-docx, matplotlib)", memo)
    check("PDF write/read/render (reportlab, pdfplumber, pypdfium2)", pdf_roundtrip)
    check(f"{SECRET_STORE_LABEL} backend", keychain)
    check("external tools", external_tools)

    shutil.rmtree(scratch, ignore_errors=True)
    print("SELF-TEST " + ("FAILED" if failures else "PASSED"), flush=True)
    return 1 if failures else 0
