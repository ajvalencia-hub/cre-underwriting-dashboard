"""Mapping preview: resolution, skip reasons and unit warnings mirror what
excel_writer.inject_values would actually do — without writing anything."""

from pathlib import Path

import openpyxl
import pytest
from openpyxl.workbook.defined_name import DefinedName

from app.services import excel_writer, mapping_preview


@pytest.fixture
def template(tmp_path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inputs"
    ws["B2"] = 12_000_000  # purchase price
    ws["B3"] = 5.5  # exit cap typed as a whole-number percent, not % formatted
    ws["B4"] = 0.05
    ws["B4"].number_format = "0.00%"  # vacancy, percent formatted
    ws["B5"] = "=B2*0.02"  # formula
    ws["B6"] = 150_000  # monthly rent where the app holds annual
    ws["B8"] = "=B2/B6"  # output formula
    ws.merge_cells("D2:E2")
    wb.defined_names["PurchasePrice"] = DefinedName("PurchasePrice", attr_text="Inputs!$B$2")
    path = tmp_path / "t.xlsx"
    wb.save(path)
    return path


def _by_id(rows):
    return {r["fieldId"]: r for r in rows}


def test_preview_statuses_match_injection(template, tmp_path):
    mappings = {
        "purchasePrice": {"target": "namedRange", "ref": "PurchasePrice"},
        "exitCapRatePct": {"target": "cell", "ref": "Inputs!B3", "sheet": "Inputs"},
        "vacancyPct": {"target": "cell", "ref": "Inputs!B4", "sheet": "Inputs"},
        "closingCostsPct": {"target": "cell", "ref": "Inputs!B5", "sheet": "Inputs"},
        "grossPotentialRent": {"target": "cell", "ref": "Inputs!B6", "sheet": "Inputs"},
        "holdPeriodYears": {"target": "cell", "ref": "Inputs!E2", "sheet": "Inputs"},
        "dealName": {"target": "cell", "ref": "Missing!A1", "sheet": "Missing"},
        "leveredIrr": {"target": "cell", "ref": "Inputs!B8", "sheet": "Inputs"},
    }
    values = {
        "purchasePrice": 12_500_000,
        "exitCapRatePct": 0.055,
        "vacancyPct": 0.06,
        "closingCostsPct": 0.02,
        "grossPotentialRent": 1_800_000,
        "holdPeriodYears": None,
        "dealName": "Test",
        "interestRate": 0.065,  # has a value, not mapped
    }
    rows = _by_id(mapping_preview.preview(template, mappings, values))

    assert rows["purchasePrice"]["resolvedRef"] == "Inputs!B2"
    assert rows["purchasePrice"]["status"] == "ok"
    assert rows["exitCapRatePct"]["status"] == "unitWarning"
    assert "whole-number percents" in rows["exitCapRatePct"]["message"]
    assert rows["vacancyPct"]["status"] == "ok"
    assert rows["closingCostsPct"]["status"] == "formula"
    assert rows["grossPotentialRent"]["status"] == "unitWarning"
    assert "monthly" in rows["grossPotentialRent"]["message"]
    assert rows["holdPeriodYears"]["status"] == "blank"
    assert rows["holdPeriodYears"]["resolvedRef"] == "Inputs!D2"  # merged -> anchor
    assert rows["dealName"]["status"] == "unresolved"
    assert rows["interestRate"]["status"] == "unmapped"
    assert "isn't mapped" in rows["interestRate"]["message"]
    assert rows["leveredIrr"]["status"] == "output" and rows["leveredIrr"]["isFormula"]

    # Cross-check against the real writer: exactly the "ok"/"unitWarning"
    # rows are written; formula / blank / unresolved rows are not.
    out = tmp_path / "out.xlsx"
    written = set(excel_writer.inject_values(template, out, mappings, values)["written"])
    expected = {fid for fid, r in rows.items() if r["status"] in ("ok", "unitWarning") and not r["isOutput"]}
    assert written == expected


def test_preview_never_modifies_the_template(template):
    before = template.read_bytes()
    mapping_preview.preview(template, {"purchasePrice": {"target": "cell", "ref": "Inputs!B2", "sheet": "Inputs"}}, {"purchasePrice": 1})
    assert template.read_bytes() == before


@pytest.mark.parametrize(
    ("field_type", "new", "old", "fmt", "warns"),
    [
        ("percent", 0.055, 5.5, "General", True),
        ("percent", 0.055, 0.05, "0.00%", False),
        ("percent", 0.055, 0.05, "General", False),
        ("currency", 1_800_000, 150_000, "General", True),  # 12x
        ("currency", 1_250_000, 1_200_000, "General", False),
        ("currency", 12_500, 12_500_000, "General", True),  # thousands
        ("currency", 100, 0, "General", False),
    ],
)
def test_unit_warning(field_type, new, old, fmt, warns):
    assert (mapping_preview.unit_warning(field_type, new, old, fmt) is not None) is warns
