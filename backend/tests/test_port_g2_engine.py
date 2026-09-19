"""Run 6 engine fixes ported onto later-items (G2): B1 sale-at-stabilization
takes no perm loan, B2 legacy stabilized NOI honors reassessed taxes, B3
cash-on-cash strips capital events (junior payoff, escrow release, maturity
refinance), B5 no-construction warning, B7 opt-in flag, validation warnings,
junior tranche on a no-takeout development, building RSF, the renovation
double display, and the additive irrDiagnostics block."""

import json
from pathlib import Path

import pytest

from app.services import excel_model_export
from app.services.proforma import engine, hold, operations, returns

FIXTURES = Path(__file__).parent / "fixtures"
REG_FIXTURES = Path(__file__).parent / "regression" / "fixtures"
CORPUS = Path(__file__).parent / "parity" / "corpus"


@pytest.fixture
def acquisition() -> dict:
    return json.loads((FIXTURES / "analytic_acquisition.json").read_text())


@pytest.fixture
def development() -> dict:
    return json.loads((FIXTURES / "analytic_development.json").read_text())


@pytest.fixture
def rollover() -> dict:
    """Detail-mode lease deal with an insurance line and debt (H3 stress on)."""
    return json.loads((REG_FIXTURES / "commercial_rollover.json").read_text())


@pytest.fixture
def feature_on() -> dict:
    return json.loads((REG_FIXTURES / "feature_on_value_add.json").read_text())


def _nnn(**overrides) -> dict:
    deal = json.loads((CORPUS / "commercial_nnn" / "inputs.json").read_text())
    deal.update(overrides)
    return deal


# ---------------------------------------------------------------------------
# B1 — a development sold IN its stabilization month never takes out.
# ---------------------------------------------------------------------------
def _sale_at_stabilization(development: dict, **overrides) -> dict:
    # 18 mo construction + 12 mo lease-up -> stabilization month 31.
    return {**development, "holdPeriodYears": 31 / 12, **overrides}


def test_b1_sold_in_stabilization_month_has_no_takeout(development):
    result = engine.compute(_sale_at_stabilization(development, refiCostsPct=0.03))
    stmt = result["statement"]
    exit_month = stmt["exitMonth"]
    assert exit_month == stmt["stabilizationMonth"] == 31
    assert stmt["loanFees"][exit_month] == 0.0
    assert stmt["debtDraws"][exit_month] == 0.0
    assert any("repaid from sale proceeds" in w for w in result["warnings"])
    assert stmt["saleProceedsNet"][exit_month] == pytest.approx(
        stmt["saleProceedsGross"][exit_month] - stmt["loanBalance"][exit_month], abs=1e-6
    )


def test_b1_sale_leg_is_independent_of_refi_costs(development):
    free = hold.refi_vs_sale({**development, "refiCostsPct": 0.0})
    costly = hold.refi_vs_sale({**development, "refiCostsPct": 0.03})
    assert free["saleAtStabilization"]["leveredIrr"] == pytest.approx(
        costly["saleAtStabilization"]["leveredIrr"], abs=1e-12
    )
    assert free["holdThroughRefi"]["leveredIrr"] > costly["holdThroughRefi"]["leveredIrr"]


def test_b1_takeout_still_fires_when_exit_is_after_stabilization(development):
    result = engine.compute({**development, "refiCostsPct": 0.03})
    stab = result["statement"]["stabilizationMonth"]
    assert result["statement"]["loanFees"][stab] > 0


def test_b1_export_refuses_sale_in_stabilization_month(development):
    deal = _sale_at_stabilization({**development, "waterfallTiers": []})
    blockers = excel_model_export.unsupported_features(deal)
    assert any("stabilization month" in b for b in blockers)
    # One month later the takeout fires and the shape exports as before.
    later = {**deal, "holdPeriodYears": 32 / 12}
    assert not any("stabiliz" in b for b in excel_model_export.unsupported_features(later))


def test_prepayment_cost_is_mirrored_in_the_export(acquisition):
    """The Excel refusal is lifted: net sale proceeds (Outputs!B8) and the
    Model exit flow carry balance x (1 + pct). Parity cases
    export_prepayment_{acquisition,development} recalc it in LibreOffice."""
    from io import BytesIO

    import openpyxl

    deal = {**acquisition, "prepaymentPenaltyPct": 0.02}
    assert engine.compute(deal)["outputs"]["prepaymentCost"] > 0
    data, _ = excel_model_export.build_model_workbook(deal)
    wb = openpyxl.load_workbook(BytesIO(data))
    inputs_ws = wb["Inputs"]
    rows = {inputs_ws.cell(row=r, column=1).value: r for r in range(1, 80)}
    prepay_row = rows["Prepayment cost % (of loan balance repaid at sale)"]
    assert inputs_ws.cell(row=prepay_row, column=2).value == 0.02
    assert f"$B${prepay_row}" in wb["Outputs"]["B8"].value
    hold_row = int(acquisition["holdPeriodYears"] * 12) + 1
    assert f"$B${prepay_row}" in wb["Model"].cell(row=hold_row, column=19).value


# ---------------------------------------------------------------------------
# B2 — legacy stabilized NOI honors useReassessedTaxes.
# ---------------------------------------------------------------------------
def test_b2_reassessed_taxes_move_legacy_stabilized_noi(acquisition):
    base_inputs = {**acquisition, "loanAmount": 0, "ltvOrLtc": 0.65}
    reassessed_inputs = {
        **base_inputs, "useReassessedTaxes": True,
        "millageRatePct": 0.02, "assessmentRatio": 0.85,
    }
    assert not operations.has_opex_detail(base_inputs)
    delta = base_inputs["purchasePrice"] * 0.85 * 0.02 - base_inputs["realEstateTaxes"]
    assert delta != 0
    base_noi = operations.stabilized_annual_noi(base_inputs)
    reassessed_noi = operations.stabilized_annual_noi(reassessed_inputs)
    assert base_noi - reassessed_noi == pytest.approx(delta, abs=1e-6)
    reassessed = engine.compute(reassessed_inputs)
    base = engine.compute(base_inputs)
    assert reassessed["debt"]["sizingNoi"] == pytest.approx(base["debt"]["sizingNoi"] - delta)
    # The vector-based year-1 NOI and the legacy figure now agree on the tax.
    assert sum(reassessed["statement"]["noi"][1:13]) == pytest.approx(reassessed_noi, rel=1e-9)


def test_b2_toggle_off_is_identical(acquisition):
    off = {**acquisition, "useReassessedTaxes": False, "millageRatePct": 0.02}
    assert engine.compute(acquisition)["outputs"] == engine.compute(off)["outputs"]


# ---------------------------------------------------------------------------
# B3 — cash-on-cash strips every capital event.
# ---------------------------------------------------------------------------
def _flat(acquisition: dict) -> dict:
    return {**acquisition, "rentGrowthMode": "flat", "expenseGrowthMode": "flat",
            "ioMonths": 60, "holdPeriodYears": 5}


def test_b3_mezz_payoff_is_stripped(acquisition):
    deal = {**_flat(acquisition), "juniorTrancheKind": "mezz", "juniorAmount": 100000,
            "juniorRatePct": 0.12, "juniorPayMode": "current"}
    result = engine.compute(deal)
    assert result["juniorTranche"]["payoff"] > 0
    out = result["outputs"]
    assert out["avgCashOnCash"] == pytest.approx(out["cashOnCashYear1"], rel=1e-9)


def test_b3_escrow_release_is_stripped(acquisition):
    result = engine.compute({**_flat(acquisition), "monthsOfTaxesAndInsurance": 6})
    stmt = result["statement"]
    assert stmt["escrowFlows"][stmt["exitMonth"]] > 0
    out = result["outputs"]
    assert out["avgCashOnCash"] == pytest.approx(out["cashOnCashYear1"], rel=1e-9)


def test_b3_maturity_refinance_flow_is_stripped(acquisition):
    """The target's own maturity refinance (roadmap #11) puts the net
    cash-out / paydown in the maturity month — a capital event, not
    operating cash."""
    deal = {**_flat(acquisition), "loanTermYears": 3, "refiCostsPct": 0.01}
    result = engine.compute(deal)
    refi = result["maturityRefinance"]
    assert refi is not None and refi["netToEquity"] != 0
    stmt = result["statement"]
    total = stmt["exitMonth"]
    operating = stmt["levered"][1 : total + 1]
    operating[-1] -= stmt["saleProceedsNet"][total]
    operating[refi["month"] - 1] -= refi["netToEquity"]
    equity_in = -sum(cf for cf in stmt["levered"] if cf < 0)
    yearly = [sum(operating[y * 12 : (y + 1) * 12]) for y in range(total // 12)]
    assert result["outputs"]["avgCashOnCash"] == pytest.approx(
        sum(yearly) / len(yearly) / equity_in, rel=1e-9
    )


def test_b3_insurance_stress_delta_ignores_escrow_release(rollover):
    """The H3 helper strips the same capital events: an escrow (which grows
    with the bumped insurance line) must not leak into the operating delta."""
    plain = engine.compute(rollover)["debt"]["insuranceStress"]
    escrowed = engine.compute({**rollover, "monthsOfTaxesAndInsurance": 6})["debt"]["insuranceStress"]
    for a, b in zip(plain, escrowed):
        assert b["leveredCfDeltaAnnual"] == pytest.approx(a["leveredCfDeltaAnnual"], rel=1e-9)


# ---------------------------------------------------------------------------
# B5 — a development with no construction period warns (engine only).
# ---------------------------------------------------------------------------
def test_b5_zero_construction_warns(development):
    result = engine.compute({**development, "constructionMonths": 0})
    assert any("no construction period" in w for w in result["warnings"])
    assert not any("no construction period" in w for w in engine.compute(development)["warnings"])


# ---------------------------------------------------------------------------
# B7 — detail-mode year-1 recoveries: implemented, OFF.
# ---------------------------------------------------------------------------
def test_b7_flag_defaults_off_and_opts_in(rollover, monkeypatch):
    assert operations.DETAIL_MODE_YEAR1_RECOVERIES is False
    _, other_flat, source, _ = operations.annual_gpr_and_other_income(rollover)
    assert source == "commercialLeases"
    before = engine.compute(rollover)["outputs"]
    monkeypatch.setattr(operations, "DETAIL_MODE_YEAR1_RECOVERIES", True)
    _, other_detail, _, _ = operations.annual_gpr_and_other_income(rollover)
    assert other_detail > other_flat
    after = engine.compute(rollover)["outputs"]
    assert after["breakEvenRatio"] > before["breakEvenRatio"]
    # EGI and opex rise by the same recovery amount: occupancy is invariant.
    assert after["breakEvenOccupancy"] == pytest.approx(before["breakEvenOccupancy"], abs=1e-9)
    for key, value in before.items():
        if key != "breakEvenRatio" and isinstance(value, float):
            assert after[key] == pytest.approx(value, abs=1e-9), key


# ---------------------------------------------------------------------------
# Validation warnings (never errors).
# ---------------------------------------------------------------------------
_V_MARKERS = ("double counting", "charged twice", "below 1%", "in_place on a development")


def test_validation_warnings(acquisition, development):
    acq = {
        **acquisition,
        "unitMix": [{"unitType": "1BR", "unitCount": 10, "inPlaceRent": 1000,
                     "marketRent": 1100, "annualTurnoverPct": 0.5}],
        "lossToLeasePct": 0.05,
        "replacementReserves": 10000, "replacementReservesPerUnit": 300,
        "exitCapRatePct": 0.005,
    }
    warnings = engine.compute(acq)["warnings"]
    for marker in _V_MARKERS[:3]:
        assert any(marker in w for w in warnings), marker
    dev = engine.compute({**development, "sizingNoiBasis": "in_place"})["warnings"]
    assert any("in_place on a development" in w for w in dev)
    # An explicit in-place NOI is what the basis is for: no warning.
    dev_ok = engine.compute({**development, "sizingNoiBasis": "in_place", "inPlaceNoi": 900000})
    assert not any("in_place on a development" in w for w in dev_ok["warnings"])


@pytest.mark.parametrize("name", [
    "analytic_acquisition", "analytic_development", "commercial_nnn",
    "commercial_rollover", "mixed_use", "value_add_multifamily", "feature_on_value_add",
])
def test_no_regression_fixture_triggers_the_new_warnings(name):
    from tests.regression.test_run4_baseline import CASES

    result = engine.compute(json.loads(CASES[name].read_text()))
    for marker in (*_V_MARKERS, "no construction period", "Junior tranche would fund"):
        assert not any(marker in w for w in result["warnings"]), (name, marker)
    assert "irrDiagnostics" not in result


# ---------------------------------------------------------------------------
# Junior tranche on a development that never takes out.
# ---------------------------------------------------------------------------
_MEZZ = {"juniorTrancheKind": "mezz", "juniorAmount": 1_000_000, "juniorRatePct": 0.12,
         "juniorOriginationFeePct": 0.01, "juniorPayMode": "current"}


@pytest.mark.parametrize("hold_years", [2, 31 / 12])
def test_junior_on_no_takeout_development_is_skipped(development, hold_years):
    plain = engine.compute({**development, "holdPeriodYears": hold_years})
    result = engine.compute({**development, "holdPeriodYears": hold_years, **_MEZZ})
    assert any("Junior tranche would fund in the exit month" in w for w in result["warnings"])
    assert result["juniorTranche"] is None
    assert "juniorPayoff" not in result["statement"]
    assert result["outputs"]["leveredIrr"] == pytest.approx(plain["outputs"]["leveredIrr"], abs=1e-12)
    assert result["outputs"]["totalProfit"] == pytest.approx(plain["outputs"]["totalProfit"], abs=1e-6)


def test_junior_with_a_real_takeout_is_unchanged(development):
    result = engine.compute({**development, **_MEZZ})
    assert result["juniorTranche"]["fundMonth"] == result["statement"]["stabilizationMonth"]
    assert not any("Junior tranche would fund" in w for w in result["warnings"])


# ---------------------------------------------------------------------------
# Renovation funded at close is displayed once (renovationCapex only).
# ---------------------------------------------------------------------------
def test_reno_equity_at_close_is_not_double_displayed(feature_on):
    stmt = engine.compute(feature_on)["statement"]
    budget = stmt["renovation"]["budget"]
    assert stmt["renovationCapex"][0] == pytest.approx(budget)
    price = feature_on["purchasePrice"]
    basis = (
        price * (1 + feature_on["closingCostsPct"] + feature_on["acquisitionFeePct"])
        + feature_on["dueDiligenceCosts"] + feature_on["dayOneCapex"]
    )
    assert stmt["costs"][0] == pytest.approx(basis)


# ---------------------------------------------------------------------------
# Building RSF (lease deals).
# ---------------------------------------------------------------------------
def test_blank_building_rsf_is_identical():
    base = engine.compute(_nnn())
    assert engine.compute(_nnn(buildingRsf=None)) == base
    assert engine.compute(_nnn(buildingRsf=0)) == base
    assert "buildingRsf" not in base["statement"]["leases"]


def test_building_rsf_above_listed_sf_reduces_recoveries_and_occupancy():
    base = engine.compute(_nnn())
    listed = sum(lease["sf"] for lease in _nnn()["commercialLeases"])
    bigger = engine.compute(_nnn(buildingRsf=listed * 1.25))
    sb, s = base["statement"], bigger["statement"]
    for m in range(1, s["exitMonth"] + 1):
        assert s["recoveries"][m] == pytest.approx(0.8 * sb["recoveries"][m])
        assert s["gpr"][m] == pytest.approx(sb["gpr"][m])
        assert s["occupancy"][m] == pytest.approx(0.8 * sb["occupancy"][m])
    assert s["leases"]["buildingRsf"] == pytest.approx(listed * 1.25)
    assert s["leases"]["totalSf"] == pytest.approx(listed)
    assert bigger["outputs"]["leveredIrr"] < base["outputs"]["leveredIrr"]


def test_building_rsf_below_listed_sf_is_ignored_with_warning():
    base = engine.compute(_nnn())
    smaller = engine.compute(_nnn(buildingRsf=100))
    assert smaller["outputs"] == base["outputs"]
    assert "buildingRsf" not in smaller["statement"]["leases"]
    assert any("buildingRsf" in w and "below the listed lease SF" in w for w in smaller["warnings"])


def test_building_rsf_with_general_vacancy_and_profiles():
    """Target features on top: general vacancy tops up against the BILLED
    potential (scheduled rent + the tenants' own recoveries — unlisted suites
    carry no modeled rent), and a lease on a market leasing profile still
    measures occupancy over the whole building."""
    listed = sum(lease["sf"] for lease in _nnn()["commercialLeases"])
    leases = [{**lease, "leasingProfile": "Retail"} for lease in _nnn()["commercialLeases"]]
    deal = _nnn(
        buildingRsf=listed * 2, leaseGeneralVacancyPct=0.05, commercialLeases=leases,
        marketLeasingProfiles=[{"profileName": "Retail", "downtimeMonths": 3}],
    )
    no_rsf = {k: v for k, v in deal.items() if k != "buildingRsf"}
    s = engine.compute(deal)["statement"]
    s0 = engine.compute(no_rsf)["statement"]
    for m in range(1, 13):
        assert s["occupancy"][m] == pytest.approx(0.5 * s0["occupancy"][m])
        expected_general = 0.05 * (s["gpr"][m] + s["recoveries"][m])
        assert s["vacancyLoss"][m] == pytest.approx(expected_general, rel=1e-9)


# ---------------------------------------------------------------------------
# irrDiagnostics — additive only; the reported IRR is never re-selected.
# ---------------------------------------------------------------------------
def test_irr_diagnostics_block_is_additive(monkeypatch, acquisition):
    base = engine.compute(acquisition)
    assert "irrDiagnostics" not in base
    monkeypatch.setattr(returns, "sign_changes", lambda flows: 3)
    monkeypatch.setattr(returns, "periodic_irr_roots", lambda flows: [-0.995, 0.05, 0.31, 4.5])
    result = engine.compute(acquisition)
    diag = result["irrDiagnostics"]
    assert diag["irrMultipleRoots"] is True
    for series in ("levered", "unlevered"):
        assert diag[series]["signChanges"] == 3
        assert diag[series]["roots"] == [0.05, 0.31]  # only the -99%..300% band
    assert diag["levered"]["reported"] == result["outputs"]["leveredIrr"]
    # The reported IRRs are untouched, and the warning still names the roots.
    assert result["outputs"]["leveredIrr"] == base["outputs"]["leveredIrr"]
    assert result["outputs"]["unleveredIrr"] == base["outputs"]["unleveredIrr"]
    assert any("change sign more than once" in w for w in result["warnings"])


def test_irr_diagnostics_absent_when_a_single_root(monkeypatch, acquisition):
    monkeypatch.setattr(returns, "sign_changes", lambda flows: 2)
    monkeypatch.setattr(returns, "periodic_irr_roots", lambda flows: [0.05])
    assert "irrDiagnostics" not in engine.compute(acquisition)
