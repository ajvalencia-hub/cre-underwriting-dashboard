"""Regression coverage for a real-world extraction failure mode: a single
sheet stacking a rent roll (no Unit Type column, "Unit N" ids, vacancy
marked only in the unit label) directly above a two-column "label: value"
income statement with duplicated CURRENT IN-PLACE vs PRO-FORMA sections.
Traced from an actual deal package where this shape caused every
income-statement row to leak into the rent roll as a phantom unit, vacant
units to read as occupied, and none of the headline financial figures
(purchase price, NOI, expenses) to reach any Deal Input field at all.

See also: tests/extraction_corpus/builders.py's
build_combined_rent_roll_and_income_statement, and
tests/test_extraction_golden.py for the broader golden-file suite this
complements (this fixture is asserted directly rather than via golden diff,
since its value is in named, readable expectations for each bug it guards).
"""

from app.models import Document
from app.services import extraction_service
from app.services.extraction import excel_extractor, operating_statement_parser, rent_roll_parser
from tests.extraction_corpus import builders


def _grid(tmp_path):
    path = tmp_path / "combined.xlsx"
    builders.build_combined_rent_roll_and_income_statement(path)
    return excel_extractor.extract_grid(path, "xlsx")


def test_rent_roll_boundary_stops_before_income_statement(tmp_path):
    """The income-statement rows below TOTAL:/AVERAGE: must never become
    phantom units, no matter how many of them have SOMETHING in the unit
    column (every "Property Taxes" / "PRO-FORMA" / etc. row would otherwise
    leak in as a fake unit)."""
    grid = _grid(tmp_path)
    parsed = rent_roll_parser.parse_rows(grid["headers"], grid["rows"], "combined.xlsx", grid["sheet"])

    assert [r["unit"] for r in parsed["rows"]] == ["Unit 1", "Unit 2", "Unit 3 - Vacant", "Unit 4"]


def test_vacant_by_unit_label_not_masked_by_placeholder_tenant(tmp_path):
    """Unit 3's tenant column says "Residential" (a boilerplate placeholder,
    not a real tenant) — only the unit LABEL says Vacant. Status must still
    read vacant, and its rent must not count toward occupied income."""
    grid = _grid(tmp_path)
    parsed = rent_roll_parser.parse_rows(grid["headers"], grid["rows"], "combined.xlsx", grid["sheet"])
    by_unit = {r["unit"]: r for r in parsed["rows"]}

    assert by_unit["Unit 3 - Vacant"]["status"] == "vacant"
    assert by_unit["Unit 1"]["status"] == "occupied"


def test_multifamily_detected_without_unit_type_column(tmp_path):
    """No Unit Type column at all — must still route to vacancyPct
    (multifamily), not retailVacancyPct (the commercial fallback), via the
    generic "Residential" tenant + "Unit N" id structural signal."""
    grid = _grid(tmp_path)
    parsed = rent_roll_parser.parse_rows(grid["headers"], grid["rows"], "combined.xlsx", grid["sheet"])

    assert extraction_service._looks_multifamily(parsed["rows"]) is True


def test_unit_mix_falls_back_to_sf_grouping(tmp_path):
    grid = _grid(tmp_path)
    parsed = rent_roll_parser.parse_rows(grid["headers"], grid["rows"], "combined.xlsx", grid["sheet"])
    proposal = rent_roll_parser.propose_unit_mix(parsed["rows"])

    assert proposal["groupedBy"] == "sf"
    assert any("square footage" in w.lower() for w in proposal["warnings"])
    by_type = {r["unitType"]: r for r in proposal["rows"]}
    assert by_type["400 SF"]["unitCount"] == 3
    assert by_type["400 SF"]["occupiedCount"] == 2  # Unit 3 is vacant
    assert by_type["600 SF"]["unitCount"] == 1


def test_operating_statement_parser_pro_forma_wins_and_expenses_sum(tmp_path):
    """The core of the fix: purchase price and both NOI scalars land on real
    field ids; the pro-forma section's figures win over the in-place
    duplicates for shared line items; Electric + Water/Sewer both rolling up
    to "utilities" are SUMMED, not last-one-wins."""
    grid = _grid(tmp_path)
    result = operating_statement_parser.parse_label_value_pairs(
        grid["headers"], grid["rows"], "combined.xlsx", grid["sheet"]
    )
    by_field = {s["fieldId"]: s["value"] for s in result["scalars"]}

    assert by_field["purchasePrice"] == 900000
    assert by_field["inPlaceNoi"] == 43700
    assert by_field["stabilizedNoi"] == 58900  # NOT the "@ 100% occupancy w/ 3rd party mgmt" variant (61000)
    assert by_field["grossPotentialRent"] == 76800  # pro-forma, not the in-place 61200
    assert by_field["realEstateTaxes"] == 8200  # pro-forma, not the in-place 8000
    assert by_field["utilities"] == 1500  # pro-forma Electric(600) + Water/Sewer(900) summed


def test_end_to_end_extraction_populates_deal_input_fields(tmp_path):
    """Through the real dispatcher: a document classified rent_roll still
    yields the income-statement scalars, merged alongside the rent-roll
    aggregation — this is the actual bug (extraction_service dispatched
    EXCLUSIVELY on document_type, so a combined sheet classified rent_roll
    never got its income-statement content parsed at all)."""
    path = tmp_path / "combined.xlsx"
    builders.build_combined_rent_roll_and_income_statement(path)
    doc = Document(
        filename="combined.xlsx", file_hash="h", stored_path=str(path), file_ext="xlsx",
        document_type="rent_roll", type_confidence=1.0,
    )
    outcome = extraction_service.run_extraction([doc])
    fields = outcome["fields"]

    assert fields["purchasePrice"]["value"] == 900000
    assert fields["inPlaceNoi"]["value"] == 43700
    assert fields["stabilizedNoi"]["value"] == 58900
    assert fields["grossPotentialRent"]["value"] == 76800
    assert fields["realEstateTaxes"]["value"] == 8200
    assert fields["utilities"]["value"] == 1500
    assert fields["vacancyPct"]["value"] == 0.25
    assert "retailVacancyPct" not in fields
    assert [r["unitType"] for r in fields["unitMix"]["value"]] == ["400 SF", "600 SF"]
    # No phantom units means no unmatched noise either.
    assert outcome["unmatchedExtractions"] == []


def test_opstmt_parser_ignores_grids_below_the_match_threshold(tmp_path):
    """A grid with only 1-2 recognizable label:value pairs (a coincidental
    match, not a real income statement) must not contribute noisy scalars —
    extraction_service gates on _MIN_OPSTMT_MATCHES."""
    path = tmp_path / "sparse.xlsx"
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Unit", "Tenant", "SF", "Rent"])
    ws.append(["A-1", "Jane Doe", 700, 1500])
    ws.append(["Insurance", 500])  # only one recognizable opstmt pair
    wb.save(path)

    doc = Document(
        filename="sparse.xlsx", file_hash="h2", stored_path=str(path), file_ext="xlsx",
        document_type="rent_roll", type_confidence=1.0,
    )
    outcome = extraction_service.run_extraction([doc])
    assert "insurance" not in outcome["fields"]
