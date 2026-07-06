"""J8: Monte Carlo — seeded determinism, sampling moments, correlation
preservation, percentile math, caps/validation, job store, memo section."""

import json
import time
from io import BytesIO
from pathlib import Path

import numpy as np
import pytest
from docx import Document

from app.services import memo_service, monte_carlo

FIXTURES = Path(__file__).parent / "fixtures"


def analytic(**overrides) -> dict:
    deal = json.loads((FIXTURES / "analytic_acquisition.json").read_text())
    deal.update(overrides)
    return deal


def _driver(path="exitCapRatePct", dist="uniform", **params):
    return {"inputPath": path, "distribution": dist,
            "params": params or {"min": 0.07, "max": 0.09}}


def test_seeded_runs_are_identical():
    kwargs = dict(
        values=analytic(), drivers=[_driver()], n=25, seed=1234, hurdle_irr=0.08
    )
    first = monte_carlo.run_simulation(**kwargs)
    second = monte_carlo.run_simulation(**kwargs)
    assert first == second
    assert first["successfulRuns"] == 25
    assert first["seed"] == 1234
    # Exit caps 7–9% straddle the 8% par cap: IRR must vary but never go
    # negative on this deal.
    assert first["probIrrNegative"] == 0
    assert first["leveredIrr"]["p5"] < first["leveredIrr"]["p95"]


def test_sampling_moments():
    n = 2000
    normal = monte_carlo.sample_matrix(
        [{"inputPath": "a", "distribution": "normal",
          "params": {"mean": 10.0, "stdDev": 2.0}}], None, n, seed=7,
    )
    values = np.array([r[0] for r in normal])
    assert values.mean() == pytest.approx(10.0, abs=0.2)
    assert values.std() == pytest.approx(2.0, abs=0.2)

    tri = monte_carlo.sample_matrix(
        [{"inputPath": "a", "distribution": "triangular",
          "params": {"min": 0.0, "mode": 1.0, "max": 2.0}}], None, n, seed=7,
    )
    tri_values = np.array([r[0] for r in tri])
    assert tri_values.mean() == pytest.approx(1.0, abs=0.06)  # (min+mode+max)/3
    assert tri_values.min() >= 0.0 and tri_values.max() <= 2.0

    uni = monte_carlo.sample_matrix(
        [{"inputPath": "a", "distribution": "uniform",
          "params": {"min": 4.0, "max": 6.0}}], None, n, seed=7,
    )
    uni_values = np.array([r[0] for r in uni])
    assert uni_values.mean() == pytest.approx(5.0, abs=0.08)
    assert uni_values.min() >= 4.0 and uni_values.max() <= 6.0


def test_correlation_is_preserved_through_the_copula():
    drivers = [
        {"inputPath": "a", "distribution": "normal", "params": {"mean": 0, "stdDev": 1}},
        {"inputPath": "b", "distribution": "uniform", "params": {"min": 0, "max": 1}},
    ]
    corr = monte_carlo._correlation_matrix(
        drivers, [{"a": "a", "b": "b", "rho": 0.9}]
    )
    samples = monte_carlo.sample_matrix(drivers, corr, 2000, seed=11)
    a = np.array([r[0] for r in samples])
    b = np.array([r[1] for r in samples])
    # Rank (Spearman) correlation survives the marginal transforms.
    rank_corr = np.corrcoef(np.argsort(np.argsort(a)), np.argsort(np.argsort(b)))[0, 1]
    assert rank_corr == pytest.approx(0.9, abs=0.05)


def test_percentiles_on_known_distribution():
    stats = monte_carlo._percentiles(sorted(float(x) for x in range(1, 101)))
    assert stats["p50"] == pytest.approx(50.5)
    assert stats["p5"] == pytest.approx(5.95)
    assert stats["p95"] == pytest.approx(95.05)
    assert stats["mean"] == pytest.approx(50.5)
    assert stats["min"] == 1.0 and stats["max"] == 100.0


def test_caps_and_validation():
    with pytest.raises(monte_carlo.MonteCarloError):
        monte_carlo.run_simulation(analytic(), [_driver()], n=2001, seed=1)
    with pytest.raises(monte_carlo.MonteCarloError):
        monte_carlo.run_simulation(
            analytic(),
            [_driver(path=p) for p in
             ("purchasePrice", "exitCapRatePct", "interestRate", "vacancyPct",
              "grossPotentialRent", "otherIncome", "insurance")],  # 7 > 6
            n=10, seed=1,
        )
    with pytest.raises(monte_carlo.MonteCarloError):
        monte_carlo._validate_drivers(
            [{"inputPath": "notAField", "distribution": "uniform",
              "params": {"min": 0, "max": 1}}]
        )
    with pytest.raises(monte_carlo.MonteCarloError):
        monte_carlo._validate_drivers(
            [{"inputPath": "vacancyPct", "distribution": "triangular",
              "params": {"min": 1, "mode": 0, "max": 2}}]  # mode < min
        )
    # Jointly inconsistent pairwise correlations -> not positive definite.
    drivers = [
        {"inputPath": p, "distribution": "uniform", "params": {"min": 0, "max": 1}}
        for p in ("a", "b", "c")
    ]
    with pytest.raises(monte_carlo.MonteCarloError):
        monte_carlo._correlation_matrix(
            drivers,
            [{"a": "a", "b": "b", "rho": 0.9}, {"a": "a", "b": "c", "rho": 0.9},
             {"a": "b", "b": "c", "rho": -0.9}],
        )


def test_job_store_polls_to_done():
    job_id = monte_carlo.start_job(
        analytic(), [_driver()], None, n=10, seed=99, hurdle_irr=0.08
    )
    deadline = time.time() + 30
    status = monte_carlo.job_status(job_id)
    while status["status"] == "running" and time.time() < deadline:
        time.sleep(0.05)
        status = monte_carlo.job_status(job_id)
    assert status["status"] == "done"
    assert status["completed"] == 10
    assert status["result"]["successfulRuns"] == 10
    assert monte_carlo.job_status("nope") is None


def test_memo_gains_risk_section_when_run_is_saved():
    result = monte_carlo.run_simulation(
        analytic(), [_driver()], n=15, seed=5, hurdle_irr=0.08
    )
    memo_bytes = memo_service.build_memo(
        deal_name="Analytic", scenario_name="Base",
        inputs=analytic(), outputs={"leveredIrr": 0.1157},
        monte_carlo=result,
    )
    doc = Document(BytesIO(memo_bytes))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "MONTE CARLO" in text.upper()  # headings render uppercased
    assert "seed 5" in text
    # Without a saved run, no risk section — never fabricated.
    plain = memo_service.build_memo(
        deal_name="Analytic", scenario_name="Base",
        inputs=analytic(), outputs={"leveredIrr": 0.1157},
    )
    plain_text = "\n".join(p.text for p in Document(BytesIO(plain)).paragraphs)
    assert "MONTE CARLO" not in plain_text.upper()
