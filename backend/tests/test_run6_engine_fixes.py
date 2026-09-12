"""Run 6 engine fixes: B1 sale-at-stabilization carries no takeout, B2 legacy
stabilized NOI honors reassessed taxes, B3 avg cash-on-cash strips exit
capital events, B4 amortYears=0 is interest-only everywhere, B5 zero
construction period warns + Excel refuses, B6 construction fee basis, B7
detail-mode year-1 recoveries (opt-in), P1 stress-skip in analysis callers,
Run-6 validation warnings, and the prepayment-penalty [FIN] feature."""

import json
from pathlib import Path

import openpyxl
import pytest

from app.services import excel_model_export, tornado_service
from app.services.proforma import debt, engine, hold, operations
from app.services.proforma.timeline import build_timeline

FIXTURES = Path(__file__).parent / "fixtures"
REG_FIXTURES = Path(__file__).parent / "regression" / "fixtures"


@pytest.fixture
def acquisition() -> dict:
    return json.loads((FIXTURES / "analytic_acquisition.json").read_text())


@pytest.fixture
def development() -> dict:
    return json.loads((FIXTURES / "analytic_development.json").read_text())


@pytest.fixture
def rollover() -> dict:
    """Detail-mode (opexLineItems) lease deal with an insurance line + debt —
    the shape that triggers the H3 insurance-stress sub-computes."""
    return json.loads((REG_FIXTURES / "commercial_rollover.json").read_text())


def _workbook(inputs: dict) -> openpyxl.Workbook:
    data, _ = excel_model_export.build_model_workbook(inputs)
    from io import BytesIO
    return openpyxl.load_workbook(BytesIO(data))


# ---------------------------------------------------------------------------
# B1 — the refi-vs-sale "sale" leg never originates the perm loan.
# ---------------------------------------------------------------------------
def test_b1_sale_leg_irr_is_independent_of_refi_costs(development):
    fork_free = hold.refi_vs_sale({**development, "refiCostsPct": 0.0})
    fork_costly = hold.refi_vs_sale({**development, "refiCostsPct": 0.03})
    assert fork_free["saleAtStabilization"] is not None
    assert fork_free["saleAtStabilization"]["leveredIrr"] == pytest.approx(
        fork_costly["saleAtStabilization"]["leveredIrr"], abs=1e-12
    )
    assert fork_free["saleAtStabilization"]["equityMultiple"] == pytest.approx(
        fork_costly["saleAtStabilization"]["equityMultiple"], abs=1e-12
    )
    # The hold-through-refi leg DOES pay the costs (sanity: the fork is live).
    assert fork_free["holdThroughRefi"]["leveredIrr"] > fork_costly["holdThroughRefi"]["leveredIrr"]


def test_b1_sale_leg_statement_has_no_takeout_in_exit_month(development):
    timeline, _ = build_timeline(
        "development", development["holdPeriodYears"],
        construction_months=development.get("constructionMonths"),
        lease_up_months=development.get("leaseUpMonths"),
        stabilization_month=development.get("stabilizationMonth"),
    )
    stab = timeline.stabilization_month
    result = engine.compute(
        {**development, "holdPeriodYears": stab / 12, "refiCostsPct": 0.03}
    )
    stmt = result["statement"]
    exit_month = stmt["exitMonth"]
    assert exit_month == stab
    assert stmt["loanFees"][exit_month] == 0.0
    assert stmt["debtDraws"][exit_month] == 0.0
    assert any("repaid from sale proceeds" in w for w in result["warnings"])
    # The construction balance is repaid from proceeds: net = gross − balance.
    assert stmt["saleProceedsNet"][exit_month] == pytest.approx(
        stmt["saleProceedsGross"][exit_month] - stmt["loanBalance"][exit_month], abs=1e-6
    )


def test_b1_takeout_still_happens_when_exit_is_after_stabilization(development):
    result = engine.compute({**development, "refiCostsPct": 0.03})
    stab = result["statement"]["stabilizationMonth"]
    assert result["statement"]["loanFees"][stab] > 0


# ---------------------------------------------------------------------------
# B2 — legacy stabilized NOI honors useReassessedTaxes.
# ---------------------------------------------------------------------------
def test_b2_reassessed_taxes_move_legacy_sizing_noi(acquisition):
    base_inputs = {**acquisition, "loanAmount": 0, "ltvOrLtc": 0.65}
    reassessed_inputs = {
        **base_inputs,
        "useReassessedTaxes": True,
        "millageRatePct": 0.02,
        "assessmentRatio": 0.85,
    }
    assert not operations.has_opex_detail(base_inputs)
    projected = base_inputs["purchasePrice"] * 0.85 * 0.02
    delta = projected - base_inputs.get("realEstateTaxes", 0.0)
    assert delta != 0

    base_noi = operations.stabilized_annual_noi(base_inputs)
    reassessed_noi = operations.stabilized_annual_noi(reassessed_inputs)
    assert base_noi - reassessed_noi == pytest.approx(delta, abs=1e-6)

    base = engine.compute(base_inputs)
    reassessed = engine.compute(reassessed_inputs)
    assert reassessed["debt"]["sizingNoi"] == pytest.approx(base["debt"]["sizingNoi"] - delta, abs=1e-6)
    assert reassessed["outputs"]["yieldOnCost"] < base["outputs"]["yieldOnCost"]
    assert reassessed["outputs"]["debtYield"] != pytest.approx(base["outputs"]["debtYield"])
    assert reassessed["outputs"]["breakEvenRatio"] > base["outputs"]["breakEvenRatio"]
    # The NOI-driven sizing candidates move (LTV governs this fixture at 65%,
    # so the governing amount itself may not).
    for constraint in ("dscr", "debtYield"):
        assert reassessed["debt"]["candidates"][constraint] < base["debt"]["candidates"][constraint]
    # The vector-based year-1 NOI and the legacy figure now agree on the tax.
    year1_noi = sum(reassessed["statement"]["noi"][1:13])
    assert year1_noi == pytest.approx(reassessed_noi, rel=1e-9)


def test_b2_toggle_off_is_identical(acquisition):
    base_inputs = {**acquisition, "loanAmount": 0, "ltvOrLtc": 0.65}
    off = {**base_inputs, "useReassessedTaxes": False, "millageRatePct": 0.02}
    assert engine.compute(base_inputs)["outputs"] == engine.compute(off)["outputs"]


# ---------------------------------------------------------------------------
# B3 — avg cash-on-cash strips the junior payoff and the escrow release.
# ---------------------------------------------------------------------------
def _flat_deal(acquisition: dict) -> dict:
    return {
        **acquisition,
        "rentGrowthMode": "flat", "rentGrowthPct": 0.0,
        "expenseGrowthMode": "flat", "expenseGrowthPct": 0.0,
        "ioMonths": 60, "holdPeriodYears": 5,
    }


def test_b3_mezz_payoff_does_not_pollute_avg_cash_on_cash(acquisition):
    deal = {
        **_flat_deal(acquisition),
        "juniorTrancheKind": "mezz", "juniorAmount": 100000,
        "juniorRatePct": 0.12, "juniorPayMode": "current",
    }
    result = engine.compute(deal)
    assert result["juniorTranche"]["payoff"] > 0
    out = result["outputs"]
    assert out["avgCashOnCash"] == pytest.approx(out["cashOnCashYear1"], rel=1e-9)


def test_b3_escrow_release_does_not_pollute_last_year_coc(acquisition):
    deal = {**_flat_deal(acquisition), "monthsOfTaxesAndInsurance": 6}
    result = engine.compute(deal)
    assert "escrowFlows" in result["statement"]
    assert result["statement"]["escrowFlows"][result["statement"]["exitMonth"]] > 0
    out = result["outputs"]
    assert out["avgCashOnCash"] == pytest.approx(out["cashOnCashYear1"], rel=1e-9)


def test_b3_plain_deal_unchanged(acquisition):
    """No tranche / escrow: the strip is a no-op (baseline-safe by inspection)."""
    result = engine.compute(_flat_deal(acquisition))
    out = result["outputs"]
    assert out["avgCashOnCash"] == pytest.approx(out["cashOnCashYear1"], rel=1e-9)


# ---------------------------------------------------------------------------
# B4 — amortYears = 0 means interest-only everywhere.
# ---------------------------------------------------------------------------
def test_b4_monthly_payment_zero_amort_is_interest():
    assert debt.monthly_payment(1_000_000, 0.06, 0) == pytest.approx(1_000_000 * 0.06 / 12)
    assert debt.annual_loan_constant(0.06, 0) == pytest.approx(0.06)
    schedule = debt.amortization_schedule(1_000_000, 0.06, 0, io_months=12, months=60)
    assert all(entry.principal == 0.0 for entry in schedule)
    assert schedule[-1].balance == pytest.approx(1_000_000)


def test_b4_engine_zero_amort_never_balloons(acquisition):
    deal = {**acquisition, "amortYears": 0, "ioMonths": 12}
    result = engine.compute(deal)
    stmt = result["statement"]
    loan = result["debt"]["loanAmount"]
    assert loan > 0
    assert all(b == pytest.approx(loan) for b in stmt["loanBalance"][1:])
    assert all(p == 0.0 for p in stmt["principal"][1:])
    out = result["outputs"]
    assert out["loanConstant"] == pytest.approx(deal["interestRate"])
    assert 0.5 < out["minDscr"] < 10
    assert out["minDscr"] == pytest.approx(out["avgDscr"], rel=1e-9) or out["minDscr"] <= out["avgDscr"]


def test_b4_excel_pmt_formula_handles_zero_amort(acquisition):
    wb = _workbook({**acquisition, "amortYears": 0, "ioMonths": 12})
    pmt = wb["Outputs"]["B1"].value
    assert "ROUND(" in pmt and "<=0" in pmt  # IO guard before any division


# ---------------------------------------------------------------------------
# B5 — constructionMonths = 0 on a development.
# ---------------------------------------------------------------------------
def test_b5_zero_construction_warns_and_excel_refuses(development):
    development = {**development, "waterfallTiers": []}  # the export refuses tiers
    deal = {**development, "constructionMonths": 0}
    result = engine.compute(deal)
    assert any("no construction period" in w for w in result["warnings"])
    blockers = excel_model_export.unsupported_features(deal)
    assert any("construction period" in b.lower() for b in blockers)
    with pytest.raises(excel_model_export.UnsupportedModelFeatures):
        excel_model_export.build_model_workbook(deal)
    # A real construction period neither warns nor refuses.
    assert not any("no construction period" in w for w in engine.compute(development)["warnings"])
    assert not excel_model_export.unsupported_features(development)


# ---------------------------------------------------------------------------
# B6 — construction origination fee basis.
# ---------------------------------------------------------------------------
def test_b6_construction_financing_fee_basis():
    schedule = [100.0, 200.0, 300.0]
    first_draw = debt.construction_financing(schedule, 150.0, 0.0, 0.01)
    assert first_draw.fee_capitalized == pytest.approx(150.0 * 0.01)  # first draw = 200-50
    commitment = debt.construction_financing(schedule, 150.0, 0.0, 0.01, fee_basis_amount=450.0)
    assert commitment.fee_capitalized == pytest.approx(450.0 * 0.01)


def test_b6_engine_default_is_first_draw_and_commitment_opts_in(development):
    assert development["originationFeePct"] > 0
    base = engine.compute(development)
    explicit_default = engine.compute({**development, "constructionFeeBasis": "first_draw"})
    assert base["outputs"] == explicit_default["outputs"]
    draws = base["statement"]["debtDraws"]
    first_draw = next(d for d in draws if d > 0)
    fee_default = dict(base["sourcesAndUses"]["uses"])["Loan fees"]
    assert fee_default == pytest.approx(first_draw * development["originationFeePct"], rel=1e-9)

    commitment = engine.compute({**development, "constructionFeeBasis": "commitment"})
    uses = dict(commitment["sourcesAndUses"]["uses"])
    total_ex_fin = sum(
        uses[k] for k in ("Land", "Hard costs", "Soft costs", "Contingency", "Developer fee")
    )
    expected_fee = total_ex_fin * development["ltvOrLtc"] * development["originationFeePct"]
    assert uses["Loan fees"] == pytest.approx(expected_fee, rel=1e-9)
    assert uses["Loan fees"] > fee_default
    assert commitment["outputs"]["leveredIrr"] < base["outputs"]["leveredIrr"]


def test_b6_excel_draws_fee_mirrors_basis(development):
    wb = _workbook({**development, "constructionFeeBasis": "commitment", "waterfallTiers": []})
    inputs_ws = wb["Inputs"]
    labels = {inputs_ws.cell(row=r, column=1).value: r for r in range(1, 60)}
    basis_row = labels["Construction fee basis (first_draw | commitment)"]
    assert inputs_ws.cell(row=basis_row, column=2).value == "commitment"
    fee_formula = wb["Draws"]["F3"].value
    assert '="commitment"' in fee_formula and f"$B${basis_row}" in fee_formula


# ---------------------------------------------------------------------------
# B7 — detail-mode year-1 recoveries (opt-in, see operations.py).
# ---------------------------------------------------------------------------
def test_b7_detail_mode_year1_recoveries_opt_in(rollover, monkeypatch):
    assert operations.has_opex_detail(rollover)
    _, other_flat, source, _ = operations.annual_gpr_and_other_income(rollover)
    assert source == "commercialLeases"
    before = engine.compute(rollover)["outputs"]

    monkeypatch.setattr(operations, "DETAIL_MODE_YEAR1_RECOVERIES", True)
    _, other_detail, _, _ = operations.annual_gpr_and_other_income(rollover)
    # The fixture has no flat recoverable fields, so the flat pool recovered
    # nothing; the detail pool recovers the flagged tax/insurance/utility lines.
    assert other_detail > other_flat
    after = engine.compute(rollover)["outputs"]
    assert after["breakEvenRatio"] == pytest.approx(0.8777284581419397, abs=1e-9)
    assert before["breakEvenRatio"] == pytest.approx(0.8428964912579363, abs=1e-9)
    # Invariant by construction: EGI and opex rise by the same recovery amount.
    assert after["breakEvenOccupancy"] == pytest.approx(before["breakEvenOccupancy"], abs=1e-9)
    # Everything outside the break-even ratio is untouched.
    for key, value in before.items():
        if key != "breakEvenRatio" and isinstance(value, float):
            assert after[key] == pytest.approx(value, abs=1e-9), key


# ---------------------------------------------------------------------------
# P1 — analysis callers skip the insurance-stress sub-computes.
# ---------------------------------------------------------------------------
def test_p1_tornado_compute_count(rollover, monkeypatch):
    plain = engine.compute(rollover)
    assert plain["debt"] and plain["debt"].get("insuranceStress")  # stress is live here

    calls = {"n": 0}
    original = engine.compute

    def counting(inputs):
        calls["n"] += 1
        return original(inputs)

    monkeypatch.setattr(engine, "compute", counting)
    engine.compute(rollover)
    assert calls["n"] == 3  # base + 2 stress bumps (nested calls are counted)

    calls["n"] = 0
    tornado_service.run_tornado(rollover, "leveredIrr")
    # 1 base + (low, high) per driver, and NO nested stress computes.
    assert calls["n"] == 2 * len(tornado_service.DRIVERS) + 1


def test_p1_hold_sweep_and_refi_fork_skip_stress(rollover, development, monkeypatch):
    calls = {"n": 0}
    original = engine.compute

    def counting(inputs):
        calls["n"] += 1
        return original(inputs)

    monkeypatch.setattr(engine, "compute", counting)
    sweep = hold.hold_sweep(rollover)
    assert calls["n"] == len(sweep["rows"])
    # Skipping the stress changes nothing the rows read.
    row = next(r for r in sweep["rows"] if r["holdYear"] == rollover["holdPeriodYears"])
    assert row["leveredIrr"] == pytest.approx(original(rollover)["outputs"]["leveredIrr"], abs=1e-12)


# ---------------------------------------------------------------------------
# Run-6 validation warnings (never errors).
# ---------------------------------------------------------------------------
def test_validation_warnings(acquisition, development):
    acq = {
        **acquisition,
        "unitMix": [{"unitType": "1BR", "unitCount": 10, "inPlaceRent": 1000,
                     "marketRent": 1100, "annualTurnoverPct": 0.5}],
        "lossToLeasePct": 0.05,
        "replacementReserves": 10000, "replacementReservesPerUnit": 300,
        "exitCapRatePct": 0.005,
    }
    warnings = engine.compute(acq)["warnings"]
    assert any("double counting" in w for w in warnings)
    assert any("charged twice" in w for w in warnings)
    assert any("below 1%" in w for w in warnings)
    dev_warnings = engine.compute({**development, "sizingNoiBasis": "in_place"})["warnings"]
    assert any("in_place on a development" in w for w in dev_warnings)
    # None of the four fire on the clean fixtures.
    for clean in (engine.compute(acquisition)["warnings"], engine.compute(development)["warnings"]):
        assert not any(
            s in w for w in clean
            for s in ("double counting", "charged twice", "below 1%", "in_place on a development")
        )


# ---------------------------------------------------------------------------
# [FIN] Prepayment penalty at exit.
# ---------------------------------------------------------------------------
def test_prepayment_penalty_default_is_identical(acquisition):
    base = engine.compute(acquisition)
    explicit = engine.compute({**acquisition, "prepaymentPenaltyPct": 0})
    assert base["outputs"] == explicit["outputs"]
    assert "prepaymentPenalty" not in base["outputs"]
    assert "prepaymentPenalty" not in base["statement"]


@pytest.mark.parametrize("fixture_name", ["acquisition", "development"])
def test_prepayment_penalty_reduces_levered_only(fixture_name, request):
    inputs = request.getfixturevalue(fixture_name)
    base = engine.compute(inputs)
    penalized = engine.compute({**inputs, "prepaymentPenaltyPct": 0.02})
    exit_month = base["statement"]["exitMonth"]
    balance = base["statement"]["loanBalance"][exit_month]
    assert balance > 0
    expected = balance * 0.02
    assert penalized["outputs"]["prepaymentPenalty"] == pytest.approx(expected, rel=1e-9)
    assert penalized["statement"]["prepaymentPenalty"][exit_month] == pytest.approx(expected, rel=1e-9)
    assert penalized["outputs"]["netSaleProceeds"] == pytest.approx(
        base["outputs"]["netSaleProceeds"] - expected, rel=1e-9
    )
    assert penalized["outputs"]["leveredIrr"] < base["outputs"]["leveredIrr"]
    assert penalized["outputs"]["equityMultiple"] < base["outputs"]["equityMultiple"]
    assert penalized["outputs"]["unleveredIrr"] == pytest.approx(base["outputs"]["unleveredIrr"], abs=1e-12)
    assert penalized["outputs"]["terminalValue"] == pytest.approx(base["outputs"]["terminalValue"], abs=1e-9)
    # Statement identity: the penalty rides inside saleProceedsNet.
    assert penalized["statement"]["saleProceedsNet"][exit_month] == pytest.approx(
        base["statement"]["saleProceedsNet"][exit_month] - expected, rel=1e-9
    )


def test_prepayment_penalty_excel_mirror(acquisition):
    wb = _workbook({**acquisition, "prepaymentPenaltyPct": 0.02})
    inputs_ws = wb["Inputs"]
    labels = {inputs_ws.cell(row=r, column=1).value: r for r in range(1, 60)}
    prepay_row = labels["Prepayment penalty % (of exit loan balance)"]
    assert inputs_ws.cell(row=prepay_row, column=2).value == 0.02
    assert f"$B${prepay_row}" in wb["Outputs"]["B8"].value
    hold_months = int(acquisition["holdPeriodYears"] * 12)
    assert f"$B${prepay_row}" in wb["Model"].cell(row=hold_months + 1, column=19).value
