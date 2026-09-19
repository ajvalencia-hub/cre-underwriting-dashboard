"""LibreOffice-vs-Excel agreement check on an unmodified template (roadmap #13)."""

from pathlib import Path

import openpyxl
import pytest

from app.services import recalc_agreement, recalc_service

pytestmark = pytest.mark.skipif(not recalc_service.is_available(), reason="LibreOffice not installed")


def _template(tmp_path: Path, cached_irr: float) -> Path:
    """A workbook with an output formula and a cached ('Excel-saved') value.
    openpyxl can't write cached values, so build it, let LibreOffice compute
    it once, then overwrite the cached value in the XML to simulate Excel
    having saved a different number."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Out"
    ws["A1"] = -100
    ws["A2"] = 60
    ws["A3"] = 60
    ws["B1"] = "=IRR(A1:A3)"
    path = tmp_path / "t.xlsx"
    wb.save(path)
    recalc_service.recalc_with_libreoffice(path)
    import zipfile, re
    src = zipfile.ZipFile(path)
    items = {n: src.read(n) for n in src.namelist()}
    src.close()
    sheet = next(n for n in items if n.startswith("xl/worksheets/sheet"))
    xml = items[sheet].decode()
    xml = re.sub(r'(<c r="B1"[^>]*>.*?<v>)([^<]*)(</v>)', lambda m: f"{m.group(1)}{cached_irr}{m.group(3)}", xml, flags=re.S)
    items[sheet] = xml.encode()
    with zipfile.ZipFile(path, "w") as out:
        for name, data in items.items():
            out.writestr(name, data)
    return path


MAPPING = {"leveredIrr": {"target": "cell", "ref": "Out!B1", "sheet": "Out"}}


def test_agreement_when_the_saved_value_matches(tmp_path):
    irr = 0.130662386291808  # IRR of -100, 60, 60
    result = recalc_agreement.check(_template(tmp_path, irr), MAPPING)
    assert result["status"] == "agrees"
    assert result["rows"][0]["libreOfficeValue"] == pytest.approx(irr, rel=1e-9)


def test_a_different_saved_value_is_reported(tmp_path):
    result = recalc_agreement.check(_template(tmp_path, 0.2), MAPPING)
    assert result["status"] == "differs"
    assert result["rows"][0]["excelValue"] == pytest.approx(0.2)
    assert result["rows"][0]["agrees"] is False


def test_nothing_to_compare_without_mapped_outputs(tmp_path):
    assert recalc_agreement.check(_template(tmp_path, 0.1), {})["status"] == "noOutputsMapped"
