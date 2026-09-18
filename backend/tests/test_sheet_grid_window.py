"""Mapping picker grid: row window + number formats."""

import openpyxl

from app.services.template_service import get_sheet_grid


def test_grid_window_and_formats(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    for r in range(1, 151):
        ws.cell(row=r, column=1, value=f"label {r}")
        ws.cell(row=r, column=2, value=r / 100)
    ws["B120"].number_format = "0.00%"
    path = tmp_path / "t.xlsx"
    wb.save(path)

    first = get_sheet_grid(path, "Model")
    assert first["startRow"] == 1 and len(first["rows"]) == 60
    assert first["rows"][0][0]["ref"] == "A1"

    window = get_sheet_grid(path, "Model", max_rows=20, start_row=110)
    assert window["startRow"] == 110
    assert window["rows"][0][0]["ref"] == "A110"
    assert len(window["rows"]) == 20
    assert window["rows"][10][1]["numberFormat"] == "0.00%"  # B120
    assert window["totalRows"] == 150

    tail = get_sheet_grid(path, "Model", max_rows=60, start_row=140)
    assert len(tail["rows"]) == 11  # rows 140..150
