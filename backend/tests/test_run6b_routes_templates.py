"""Run 6b: HTTP coverage for the template / mapping / generate / sensitivity
(template mode) routes, which previously had none. A tiny workbook is built
in-process with openpyxl (two named ranges: an input cell and a formula
output cell) so every path is exercised end-to-end through the API.
LibreOffice-dependent paths skip when soffice isn't detected."""

import io
import json
from pathlib import Path

import openpyxl
import pytest
from openpyxl.workbook.defined_name import DefinedName

from app.models import Template
from app.routers import templates as templates_router
from app.services import recalc_service

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _tiny_workbook_bytes(price: float = 100) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inputs"
    ws["A1"] = "Purchase Price"
    ws["B1"] = price
    ws["A2"] = "Doubled"
    ws["B2"] = "=B1*2"
    wb.defined_names["purchasePrice"] = DefinedName("purchasePrice", attr_text="Inputs!$B$1")
    wb.defined_names["leveredIrr"] = DefinedName("leveredIrr", attr_text="Inputs!$B$2")
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _upload_template(client, name: str = "tiny.xlsx", price: float = 100) -> dict:
    response = client.post(
        "/api/templates/upload", files={"file": (name, _tiny_workbook_bytes(price), XLSX)}
    )
    assert response.status_code == 200, response.text
    return response.json()


def _named_range_profile(client, template_id: str, name: str = "P1", with_output=False) -> dict:
    mappings = {"purchasePrice": {"target": "namedRange", "ref": "purchasePrice"}}
    if with_output:
        mappings["leveredIrr"] = {"target": "namedRange", "ref": "leveredIrr"}
    response = client.post(
        "/api/mappings",
        json={"templateId": template_id, "profileName": name, "mappings": mappings},
    )
    assert response.status_code == 200, response.text
    return response.json()


# ------------------------------------------------------------------ templates


def test_list_and_get_template(client):
    assert client.get("/api/templates").json() == []
    created = _upload_template(client)
    assert created["reused"] is False
    assert [s["name"] for s in created["sheets"]] == ["Inputs"]
    assert {nr["name"] for nr in created["namedRanges"]} == {"purchasePrice", "leveredIrr"}

    listed = client.get("/api/templates").json()
    assert [t["id"] for t in listed] == [created["id"]]

    fetched = client.get(f"/api/templates/{created['id']}")
    assert fetched.status_code == 200
    assert fetched.json()["fileHash"] == created["fileHash"]
    assert client.get("/api/templates/does-not-exist").status_code == 404


def test_sheet_grid_happy_path_and_unknown_sheet(client):
    created = _upload_template(client)
    grid = client.get(f"/api/templates/{created['id']}/sheets/Inputs/grid")
    assert grid.status_code == 200, grid.text
    body = grid.json()
    assert body["sheet"] == "Inputs"
    assert body["columns"] == ["A", "B"]
    assert body["totalRows"] == 2 and body["totalCols"] == 2
    b2 = body["rows"][1][1]
    assert b2 == {"ref": "B2", "value": "=B1*2", "isFormula": True}
    assert body["rows"][0][1]["value"] == 100

    assert client.get(f"/api/templates/{created['id']}/sheets/Nope/grid").status_code == 404
    assert client.get("/api/templates/missing/sheets/Inputs/grid").status_code == 404


def test_sheet_grid_clamps_max_rows_and_cols(client, monkeypatch):
    created = _upload_template(client)
    seen: list[dict] = []

    def spy(path, sheet_name, max_rows, max_cols):
        seen.append({"max_rows": max_rows, "max_cols": max_cols})
        return {"sheet": sheet_name, "columns": [], "rows": [], "totalRows": 0, "totalCols": 0}

    monkeypatch.setattr(templates_router.template_service, "get_sheet_grid", spy)
    url = f"/api/templates/{created['id']}/sheets/Inputs/grid"
    assert client.get(url, params={"max_rows": 99999, "max_cols": 99999}).status_code == 200
    assert client.get(url, params={"max_rows": 0, "max_cols": -5}).status_code == 200
    assert client.get(url).status_code == 200
    assert seen == [
        {"max_rows": 500, "max_cols": 100},  # clamped to the ceilings
        {"max_rows": 1, "max_cols": 1},  # clamped to the floors
        {"max_rows": 60, "max_cols": 30},  # defaults
    ]


def test_sheet_grid_missing_stored_file_is_410(client, session_factory):
    created = _upload_template(client)
    with session_factory() as db:
        stored = Path(db.get(Template, created["id"]).stored_path)
    stored.unlink()
    response = client.get(f"/api/templates/{created['id']}/sheets/Inputs/grid")
    assert response.status_code == 410
    assert "re-upload" in response.json()["detail"]


def test_delete_template_cascades_profiles_scenarios_and_deal_selection(client):
    created = _upload_template(client)
    profile = _named_range_profile(client, created["id"])
    deal = client.post("/api/deals", json={"name": "Cascade"}).json()
    selected = client.put(
        f"/api/deals/{deal['id']}",
        json={"activeTemplateId": created["id"], "activeMappingProfileId": profile["id"]},
    )
    assert selected.status_code == 200, selected.text
    assert selected.json()["activeTemplateId"] == created["id"]
    scenario = client.post(
        "/api/scenarios",
        json={
            "scenarioName": "S", "kind": "full", "dealId": deal["id"],
            "templateId": created["id"], "mappingProfileId": profile["id"], "inputs": {},
        },
    )
    assert scenario.status_code == 200, scenario.text

    deleted = client.delete(f"/api/templates/{created['id']}")
    assert deleted.status_code == 200 and deleted.json() == {"deleted": True}

    assert client.get(f"/api/templates/{created['id']}").status_code == 404
    assert client.get(f"/api/mappings/{profile['id']}").status_code == 404
    assert client.get(f"/api/scenarios/{scenario.json()['id']}").status_code == 404
    after = client.get(f"/api/deals/{deal['id']}").json()
    assert after["activeTemplateId"] is None
    assert after["activeMappingProfileId"] is None
    assert client.delete(f"/api/templates/{created['id']}").status_code == 404


# ------------------------------------------------------------------- mappings


def test_mappings_list_filters_by_template(client):
    a = _upload_template(client, "a.xlsx", price=1)
    b = _upload_template(client, "b.xlsx", price=2)
    pa = _named_range_profile(client, a["id"], "A")
    pb = _named_range_profile(client, b["id"], "B")

    everything = {p["id"] for p in client.get("/api/mappings").json()}
    assert everything == {pa["id"], pb["id"]}
    only_a = client.get("/api/mappings", params={"template_id": a["id"]}).json()
    assert [p["id"] for p in only_a] == [pa["id"]]
    assert client.get("/api/mappings", params={"template_id": "nope"}).json() == []


def test_auto_match_maps_named_ranges(client):
    created = _upload_template(client)
    response = client.get(f"/api/mappings/auto-match/{created['id']}")
    assert response.status_code == 200, response.text
    mappings = response.json()["mappings"]
    assert mappings["purchasePrice"] == {
        "target": "namedRange", "ref": "purchasePrice", "anchor": None,
        "sheet": None, "columnOrder": None, "source": "auto",
    }
    assert mappings["leveredIrr"]["target"] == "namedRange"
    assert client.get("/api/mappings/auto-match/missing").status_code == 404


def test_mapping_profile_crud_and_validation(client):
    created = _upload_template(client)
    profile = _named_range_profile(client, created["id"], "Base")
    assert profile["templateId"] == created["id"]
    assert profile["profileName"] == "Base"
    assert profile["mappings"]["purchasePrice"]["source"] == "manual"
    assert isinstance(profile["unmappedRequiredFields"], list)
    assert "purchasePrice" not in profile["unmappedRequiredFields"]

    fetched = client.get(f"/api/mappings/{profile['id']}")
    assert fetched.status_code == 200 and fetched.json()["id"] == profile["id"]

    updated = client.put(
        f"/api/mappings/{profile['id']}",
        json={
            "templateId": created["id"], "profileName": "Renamed",
            "mappings": {"purchasePrice": {"target": "cell", "ref": "Inputs!B1"}},
        },
    )
    assert updated.status_code == 200, updated.text
    assert updated.json()["profileName"] == "Renamed"
    assert updated.json()["mappings"]["purchasePrice"]["target"] == "cell"

    # Validation: unknown template on create -> 404; bad target enum -> 422.
    missing_template = client.post(
        "/api/mappings",
        json={"templateId": "nope", "profileName": "X", "mappings": {}},
    )
    assert missing_template.status_code == 404
    bad_target = client.put(
        f"/api/mappings/{profile['id']}",
        json={
            "templateId": created["id"], "profileName": "X",
            "mappings": {"purchasePrice": {"target": "sheet"}},
        },
    )
    assert bad_target.status_code == 422

    assert client.delete(f"/api/mappings/{profile['id']}").json() == {"deleted": True}
    assert client.get(f"/api/mappings/{profile['id']}").status_code == 404
    assert client.put(
        f"/api/mappings/{profile['id']}",
        json={"templateId": created["id"], "profileName": "X", "mappings": {}},
    ).status_code == 404
    assert client.delete(f"/api/mappings/{profile['id']}").status_code == 404


def test_delete_mapping_clears_deal_selection_and_scenarios(client):
    created = _upload_template(client)
    profile = _named_range_profile(client, created["id"])
    deal = client.post("/api/deals", json={"name": "D"}).json()
    client.put(
        f"/api/deals/{deal['id']}",
        json={"activeTemplateId": created["id"], "activeMappingProfileId": profile["id"]},
    )
    scenario = client.post(
        "/api/scenarios",
        json={
            "scenarioName": "S", "kind": "full", "dealId": deal["id"],
            "templateId": created["id"], "mappingProfileId": profile["id"], "inputs": {},
        },
    ).json()
    assert client.delete(f"/api/mappings/{profile['id']}").status_code == 200
    assert client.get(f"/api/scenarios/{scenario['id']}").status_code == 404
    after = client.get(f"/api/deals/{deal['id']}").json()
    assert after["activeMappingProfileId"] is None
    assert after["activeTemplateId"] == created["id"]  # template selection survives


# ------------------------------------------------------------------- generate


def test_generate_template_path_writes_mapped_value_and_headers(client):
    created = _upload_template(client)
    profile = _named_range_profile(client, created["id"])
    response = client.post(
        "/api/generate",
        json={
            "templateId": created["id"], "mappingProfileId": profile["id"],
            "values": {"purchasePrice": 123456, "unmappedField": 1},
        },
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(XLSX)
    assert 'filename="tiny.xlsx"' in response.headers["content-disposition"]
    assert response.headers["x-generation-written-count"] == "1"
    assert json.loads(response.headers["x-generation-warnings"]) == []
    assert json.loads(response.headers["x-generation-outputs"]) == {}  # no recalc requested

    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    assert wb["Inputs"]["B1"].value == 123456
    assert wb["Inputs"]["B2"].value == "=B1*2"  # formula preserved, not overwritten


def test_generate_rejects_unknown_ids_and_foreign_profile(client):
    a = _upload_template(client, "a.xlsx", price=1)
    b = _upload_template(client, "b.xlsx", price=2)
    profile_b = _named_range_profile(client, b["id"])

    assert client.post(
        "/api/generate",
        json={"templateId": "nope", "mappingProfileId": profile_b["id"], "values": {}},
    ).status_code == 404
    assert client.post(
        "/api/generate",
        json={"templateId": a["id"], "mappingProfileId": "nope", "values": {}},
    ).status_code == 404
    foreign = client.post(
        "/api/generate",
        json={"templateId": a["id"], "mappingProfileId": profile_b["id"], "values": {}},
    )
    assert foreign.status_code == 400
    assert "does not belong" in foreign.json()["detail"]
    assert client.post("/api/generate", json={"templateId": a["id"]}).status_code == 422


@pytest.mark.skipif(not recalc_service.is_available(), reason="LibreOffice not detected")
def test_generate_with_recalc_reads_back_outputs(client):
    created = _upload_template(client)
    profile = _named_range_profile(client, created["id"], with_output=True)
    response = client.post(
        "/api/generate",
        json={
            "templateId": created["id"], "mappingProfileId": profile["id"],
            "values": {"purchasePrice": 21}, "recalc": True,
        },
    )
    assert response.status_code == 200, response.text
    warnings = json.loads(response.headers["x-generation-warnings"])
    assert not any("recalc skipped" in w for w in warnings), warnings
    outputs = json.loads(response.headers["x-generation-outputs"])
    assert outputs["leveredIrr"] == pytest.approx(42)
    wb = openpyxl.load_workbook(io.BytesIO(response.content), data_only=True)
    assert wb["Inputs"]["B2"].value == pytest.approx(42)


# ---------------------------------------------------------------- sensitivity


def _sensitivity_payload(template_id, profile_id, **overrides):
    payload = {
        "mode": "template", "templateId": template_id, "mappingProfileId": profile_id,
        "baseValues": {}, "outputFieldIds": ["leveredIrr"],
        "drivers": [{"fieldId": "purchasePrice", "values": [100, 250]}],
    }
    payload.update(overrides)
    return payload


def test_sensitivity_template_mode_validation_errors(client):
    created = _upload_template(client)
    profile = _named_range_profile(client, created["id"], with_output=True)

    # Malformed body (drivers missing entirely) is a schema 422.
    malformed = client.post("/api/sensitivity", json={"mode": "template", "baseValues": {}})
    assert malformed.status_code == 422

    # Zero / three drivers -> 400 before any template lookup.
    assert client.post(
        "/api/sensitivity", json=_sensitivity_payload(created["id"], profile["id"], drivers=[])
    ).status_code == 400
    three = [{"fieldId": f, "values": [1]} for f in ("purchasePrice", "a", "b")]
    assert client.post(
        "/api/sensitivity", json=_sensitivity_payload(created["id"], profile["id"], drivers=three)
    ).status_code == 400

    # Unknown template / profile -> 404; profile from another template -> 400.
    assert client.post(
        "/api/sensitivity", json=_sensitivity_payload("nope", profile["id"])
    ).status_code == 404
    assert client.post(
        "/api/sensitivity", json=_sensitivity_payload(created["id"], "nope")
    ).status_code == 404

    # A driver that isn't mapped in the profile can't move the output -> 400.
    unmapped = client.post(
        "/api/sensitivity",
        json=_sensitivity_payload(
            created["id"], profile["id"],
            drivers=[{"fieldId": "exitCapRatePct", "values": [0.05, 0.06]}],
        ),
    )
    assert unmapped.status_code == 400
    assert "exitCapRatePct" in unmapped.json()["detail"]

    # Grid too large for template mode (each point is a LibreOffice recalc).
    too_big = client.post(
        "/api/sensitivity",
        json=_sensitivity_payload(
            created["id"], profile["id"],
            drivers=[{"fieldId": "purchasePrice", "values": list(range(31))}],
        ),
    )
    assert too_big.status_code == 400
    assert "Grid too large" in too_big.json()["detail"]


@pytest.mark.skipif(not recalc_service.is_available(), reason="LibreOffice not detected")
def test_sensitivity_template_mode_recalcs_each_grid_point(client):
    created = _upload_template(client)
    profile = _named_range_profile(client, created["id"], with_output=True)
    response = client.post(
        "/api/sensitivity", json=_sensitivity_payload(created["id"], profile["id"])
    )
    assert response.status_code == 200, response.text
    points = response.json()["points"]
    assert [p["driverValues"] for p in points] == [
        {"purchasePrice": 100.0}, {"purchasePrice": 250.0},
    ]
    assert [p["outputs"]["leveredIrr"] for p in points] == pytest.approx([200, 500])
    assert all(p["warnings"] == [] for p in points)


def test_sensitivity_template_mode_reports_missing_libreoffice(client, monkeypatch):
    created = _upload_template(client)
    profile = _named_range_profile(client, created["id"], with_output=True)
    monkeypatch.setattr(recalc_service, "is_available", lambda: False)
    response = client.post(
        "/api/sensitivity", json=_sensitivity_payload(created["id"], profile["id"])
    )
    assert response.status_code == 400
    assert "LibreOffice" in response.json()["detail"]
